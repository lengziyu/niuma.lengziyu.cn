from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
import json
import re
from tempfile import TemporaryDirectory
from threading import Lock
from typing import Dict, List, Optional, Union
from urllib.parse import quote
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from pydantic import BaseModel, Field
from fastapi.responses import JSONResponse, Response
from starlette.concurrency import run_in_threadpool


def ensure_pymupdf_rect_compat() -> None:
    original_get_area = getattr(fitz.Rect, "get_area", None)

    def get_area(self, unit: str = "px") -> float:
        if original_get_area is not None and unit != "pt":
            try:
                return original_get_area(self, unit)
            except KeyError:
                pass

        width = max(0.0, float(self.x1) - float(self.x0))
        height = max(0.0, float(self.y1) - float(self.y0))
        area = width * height
        factors = {
            "px": 1.0,
            "pt": 1.0,
            "in": 1 / (72 * 72),
            "cm": (2.54 / 72) ** 2,
            "mm": (25.4 / 72) ** 2
        }
        return area * factors.get(unit, 1.0)

    fitz.Rect.get_area = get_area


ensure_pymupdf_rect_compat()

from pdf2docx import Converter


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_UPLOAD_BYTES = 80 * 1024 * 1024
MODE_EDITABLE = "editable_open_source"
MODE_VISUAL = "visual_exact"
SUPPORTED_MODES = {MODE_EDITABLE, MODE_VISUAL}
RENDER_ZOOM = 2.2
DOCX_TEXT_FONT = "Microsoft YaHei"
DOCX_EAST_ASIA_FONT = "微软雅黑"
RFONTS_RE = re.compile(rb"<w:rFonts\b[^>]*/>")
LANG_RE = re.compile(rb"<w:lang\b[^>]*/>")
TEA_GIFT_TTL = timedelta(hours=24)
TEA_GIFT_MAX_QUEUE = 20
TEA_FRIEND_ID_TTL = timedelta(days=3)
PREFERRED_SHORT_IDS = [
    "666", "888", "999", "520", "521", "233", "234", "345", "567", "678", "789", "123"
]

app = FastAPI(title="Niuma conversion service")
tea_gift_inboxes: dict[str, list[dict[str, str]]] = defaultdict(list)
tea_gift_receipts: dict[str, list[dict[str, str]]] = defaultdict(list)
tea_friend_registry: Dict[str, Dict[str, Union[str, int]]] = {}
tea_gift_lock = Lock()


@app.exception_handler(RequestValidationError)
async def request_validation_exception_handler(request: Request, exc: RequestValidationError) -> JSONResponse:
    if request.url.path.startswith("/api/tea-gifts/"):
        return JSONResponse(
            status_code=422,
            content={"detail": "奶茶请求参数不完整，请刷新页面后重试。"}
        )

    first_error = exc.errors()[0] if exc.errors() else {}
    message = first_error.get("msg") or "请求参数无效。"
    return JSONResponse(status_code=422, content={"detail": f"请求参数无效：{message}"})


class TeaGiftRequest(BaseModel):
    senderId: str = Field(min_length=3, max_length=3)
    recipientId: str = Field(min_length=3, max_length=3)
    teaKey: str = Field(min_length=1, max_length=32)
    drinkName: str = Field(min_length=1, max_length=64)
    sugar: str = Field(min_length=1, max_length=32)
    temp: str = Field(min_length=1, max_length=32)
    size: str = Field(min_length=1, max_length=32)


class TeaFriendRegisterRequest(BaseModel):
    clientToken: str = Field(min_length=8, max_length=80)
    preferredId: Optional[str] = Field(default=None, max_length=3)


def normalize_tea_gift_id(raw_value: str, field_label: str) -> str:
    normalized = re.sub(r"\D", "", (raw_value or "").strip())

    if len(normalized) != 3:
        raise HTTPException(status_code=400, detail=f"{field_label} 无效，请检查后重试。")

    return normalized


def cleanup_expired_tea_gifts(now: datetime) -> None:
    cutoff = int((now - TEA_GIFT_TTL).timestamp())
    friend_cutoff = int((now - TEA_FRIEND_ID_TTL).timestamp())
    expired_recipient_ids = []
    expired_friend_ids = []

    for recipient_id, queue in tea_gift_inboxes.items():
        fresh_queue = [gift for gift in queue if int(gift["createdAt"]) >= cutoff]

        if fresh_queue:
            tea_gift_inboxes[recipient_id] = fresh_queue
        else:
            expired_recipient_ids.append(recipient_id)

    for recipient_id in expired_recipient_ids:
        tea_gift_inboxes.pop(recipient_id, None)

    expired_receipt_ids = []
    for friend_id, queue in tea_gift_receipts.items():
        fresh_queue = [receipt for receipt in queue if int(receipt["createdAt"]) >= cutoff]

        if fresh_queue:
            tea_gift_receipts[friend_id] = fresh_queue
        else:
            expired_receipt_ids.append(friend_id)

    for friend_id in expired_receipt_ids:
        tea_gift_receipts.pop(friend_id, None)

    for friend_id, record in tea_friend_registry.items():
        if int(record["claimedAt"]) < friend_cutoff:
            expired_friend_ids.append(friend_id)

    for friend_id in expired_friend_ids:
        tea_friend_registry.pop(friend_id, None)


def normalize_client_token(raw_value: str) -> str:
    normalized = (raw_value or "").strip()
    if len(normalized) < 8:
        raise HTTPException(status_code=400, detail="设备标识无效，请刷新后重试。")
    return normalized[:80]


def is_friend_id_available(friend_id: str, client_token: str) -> bool:
    record = tea_friend_registry.get(friend_id)
    return record is None or record["clientToken"] == client_token


def allocate_short_friend_id(client_token: str) -> str:
    for preferred_id in PREFERRED_SHORT_IDS:
        if is_friend_id_available(preferred_id, client_token):
            return preferred_id

    for numeric_id in range(100, 1000):
        candidate = str(numeric_id)
        if is_friend_id_available(candidate, client_token):
            return candidate

    raise HTTPException(status_code=503, detail="在线好友编号已满，请稍后再试。")


def safe_docx_name(filename: str) -> str:
    stem = Path(filename or "converted").stem.strip() or "converted"
    return f"{stem}.docx"


def safe_sheet_title(raw_title: str, used: set[str]) -> str:
    title = re.sub(r"[:\\/?*\[\]]", "_", raw_title or "Sheet").strip() or "Sheet"
    title = title[:31]
    candidate = title
    index = 2

    while candidate in used:
        suffix = f"_{index}"
        candidate = f"{title[:31 - len(suffix)]}{suffix}"
        index += 1

    used.add(candidate)
    return candidate


def parse_pdf_page_order(page_spec: str, total_pages: int) -> list[int]:
    normalized = (page_spec or "").strip()
    if not normalized:
        return list(range(total_pages))

    pages: list[int] = []
    for part in re.split(r"[,，\s]+", normalized):
        if not part:
            continue

        match = re.fullmatch(r"(\d+)(?:-(\d+))?", part)
        if not match:
            raise HTTPException(status_code=400, detail="页码格式无效，请使用类似 1-3,5,4 的格式。")

        start = int(match.group(1))
        end = int(match.group(2) or start)
        step = 1 if start <= end else -1

        for page_number in range(start, end + step, step):
            if page_number < 1 or page_number > total_pages:
                raise HTTPException(status_code=400, detail=f"页码 {page_number} 超出范围，当前 PDF 共 {total_pages} 页。")
            pages.append(page_number - 1)

    if not pages:
        raise HTTPException(status_code=400, detail="请至少保留 1 页。")

    return pages


async def save_upload(upload: UploadFile, target: Path) -> None:
    size = 0

    with target.open("wb") as output:
        while True:
            chunk = await upload.read(1024 * 1024)
            if not chunk:
                break

            size += len(chunk)
            if size > MAX_UPLOAD_BYTES:
                raise HTTPException(status_code=413, detail="文件超过 80MB，请拆分后再转换。")

            output.write(chunk)

    if size == 0:
        raise HTTPException(status_code=400, detail="上传文件为空。")


def convert_editable_open_source(input_path: Path, output_path: Path) -> None:
    """Use LibreOffice for high-quality PDF to DOCX conversion, fall back to pdf2docx."""
    import shutil
    import subprocess

    soffice = shutil.which("soffice") or shutil.which("libreoffice")

    if soffice:
        # LibreOffice converts PDF -> DOCX with good layout preservation
        try:
            subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--norestore",
                    "--convert-to", "docx:MS Word 2007 XML",
                    "--outdir", str(output_path.parent),
                    str(input_path),
                ],
                capture_output=True,
                timeout=180,
                check=False,
            )
        except subprocess.TimeoutExpired:
            pass

        # LibreOffice outputs with the input filename stem
        lo_output = output_path.parent / f"{input_path.stem}.docx"
        if lo_output.exists() and lo_output.stat().st_size > 0:
            if lo_output != output_path:
                lo_output.rename(output_path)
            normalize_docx_text_fonts(output_path)
            return

    # Fallback to pdf2docx if LibreOffice is not available or failed
    converter = Converter(str(input_path))
    try:
        converter.convert(str(output_path), start=0, end=None, multi_processing=False)
    finally:
        converter.close()

    normalize_docx_text_fonts(output_path)


def normalize_word_xml_fonts(content: bytes) -> bytes:
    font_decl = (
        f'<w:rFonts w:ascii="{DOCX_TEXT_FONT}" '
        f'w:hAnsi="{DOCX_TEXT_FONT}" '
        f'w:eastAsia="{DOCX_EAST_ASIA_FONT}" '
        f'w:cs="{DOCX_TEXT_FONT}"/>'
    ).encode("utf-8")
    lang_decl = b'<w:lang w:val="zh-CN" w:eastAsia="zh-CN" w:bidi="zh-CN"/>'

    content = RFONTS_RE.sub(font_decl, content)
    content = LANG_RE.sub(lang_decl, content)
    return content.replace(b"MicrosoftYaHei", DOCX_TEXT_FONT.encode("utf-8"))


def normalize_docx_text_fonts(docx_path: Path) -> None:
    normalized_path = docx_path.with_suffix(".normalized.docx")

    with ZipFile(docx_path, "r") as source, ZipFile(normalized_path, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)

            if (
                item.filename.startswith("word/")
                and item.filename.endswith(".xml")
                and not item.filename.startswith("word/_rels/")
            ):
                content = normalize_word_xml_fonts(content)

            target.writestr(item, content)

    normalized_path.replace(docx_path)


def configure_section(section, width_pt: float, height_pt: float) -> None:
    section.page_width = Pt(width_pt)
    section.page_height = Pt(height_pt)
    section.top_margin = Pt(0)
    section.bottom_margin = Pt(0)
    section.left_margin = Pt(0)
    section.right_margin = Pt(0)
    section.header_distance = Pt(0)
    section.footer_distance = Pt(0)


def convert_visual_exact(input_path: Path, output_path: Path, workspace: Path) -> None:
    pdf = fitz.open(str(input_path))
    try:
        if pdf.page_count == 0:
            raise ValueError("PDF 没有可转换的页面。")

        document = Document()

        for index in range(pdf.page_count):
            page = pdf.load_page(index)
            section = document.sections[0] if index == 0 else document.add_section(WD_SECTION.NEW_PAGE)
            configure_section(section, page.rect.width, page.rect.height)

            pixmap = page.get_pixmap(matrix=fitz.Matrix(RENDER_ZOOM, RENDER_ZOOM), alpha=False)
            image_path = workspace / f"page-{index + 1}.png"
            pixmap.save(str(image_path))

            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.add_run().add_picture(str(image_path), width=section.page_width)

        document.save(str(output_path))
    finally:
        pdf.close()


def convert_pdf(input_path: Path, output_path: Path, mode: str, workspace: Path) -> None:
    if mode == MODE_VISUAL:
        convert_visual_exact(input_path, output_path, workspace)
        return

    convert_editable_open_source(input_path, output_path)


@app.get("/api/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/api/tea-gifts/register")
def register_tea_friend(payload: TeaFriendRegisterRequest) -> dict[str, str]:
    client_token = normalize_client_token(payload.clientToken)
    preferred_id = None

    if payload.preferredId:
        preferred_id = normalize_tea_gift_id(payload.preferredId, "好友 ID")

    now = datetime.now(timezone.utc)

    with tea_gift_lock:
        cleanup_expired_tea_gifts(now)

        for friend_id, record in tea_friend_registry.items():
            if record["clientToken"] == client_token:
                tea_friend_registry[friend_id]["claimedAt"] = int(now.timestamp())
                return {"friendId": friend_id}

        if preferred_id and is_friend_id_available(preferred_id, client_token):
            next_id = preferred_id
        else:
            next_id = allocate_short_friend_id(client_token)

        tea_friend_registry[next_id] = {
            "clientToken": client_token,
            "claimedAt": int(now.timestamp())
        }

    return {"friendId": next_id}


@app.post("/api/tea-gifts/send")
def send_tea_gift(payload: TeaGiftRequest) -> Dict[str, Union[str, int]]:
    sender_id = normalize_tea_gift_id(payload.senderId, "发送者 ID")
    recipient_id = normalize_tea_gift_id(payload.recipientId, "好友 ID")

    if sender_id == recipient_id:
        raise HTTPException(status_code=400, detail="不能给自己送奶茶。")

    now = datetime.now(timezone.utc)
    gift = {
        "id": uuid4().hex,
        "senderId": sender_id,
        "recipientId": recipient_id,
        "teaKey": payload.teaKey.strip(),
        "drinkName": payload.drinkName.strip(),
        "sugar": payload.sugar.strip(),
        "temp": payload.temp.strip(),
        "size": payload.size.strip(),
        "createdAt": str(int(now.timestamp()))
    }

    with tea_gift_lock:
        cleanup_expired_tea_gifts(now)
        queue = tea_gift_inboxes[recipient_id]
        queue.append(gift)
        if len(queue) > TEA_GIFT_MAX_QUEUE:
            tea_gift_inboxes[recipient_id] = queue[-TEA_GIFT_MAX_QUEUE:]
        queue_size = len(tea_gift_inboxes[recipient_id])

    return {
        "status": "queued",
        "giftId": gift["id"],
        "recipientId": recipient_id,
        "queueSize": queue_size
    }


@app.get("/api/tea-gifts/inbox/{recipient_id}")
def get_tea_gift_inbox(recipient_id: str) -> Dict[str, Union[List[Dict[str, str]], int]]:
    normalized_recipient_id = normalize_tea_gift_id(recipient_id, "好友 ID")

    with tea_gift_lock:
        cleanup_expired_tea_gifts(datetime.now(timezone.utc))
        items = list(tea_gift_inboxes.get(normalized_recipient_id, []))

    return {
        "count": len(items),
        "items": items
    }


@app.delete("/api/tea-gifts/inbox/{recipient_id}/{gift_id}")
def acknowledge_tea_gift(recipient_id: str, gift_id: str) -> Dict[str, Union[str, int]]:
    normalized_recipient_id = normalize_tea_gift_id(recipient_id, "好友 ID")
    normalized_gift_id = gift_id.strip()

    if not normalized_gift_id:
        raise HTTPException(status_code=400, detail="无效的奶茶礼物。")

    with tea_gift_lock:
        cleanup_expired_tea_gifts(datetime.now(timezone.utc))
        queue = tea_gift_inboxes.get(normalized_recipient_id, [])
        accepted_gift = next((gift for gift in queue if gift["id"] == normalized_gift_id), None)
        next_queue = [gift for gift in queue if gift["id"] != normalized_gift_id]

        if next_queue:
            tea_gift_inboxes[normalized_recipient_id] = next_queue
        else:
            tea_gift_inboxes.pop(normalized_recipient_id, None)

        if accepted_gift:
            receipt = {
                "id": uuid4().hex,
                "giftId": accepted_gift["id"],
                "senderId": accepted_gift["senderId"],
                "recipientId": normalized_recipient_id,
                "teaKey": accepted_gift["teaKey"],
                "drinkName": accepted_gift["drinkName"],
                "createdAt": str(int(datetime.now(timezone.utc).timestamp()))
            }
            receipt_queue = tea_gift_receipts[accepted_gift["senderId"]]
            receipt_queue.append(receipt)
            if len(receipt_queue) > TEA_GIFT_MAX_QUEUE:
                tea_gift_receipts[accepted_gift["senderId"]] = receipt_queue[-TEA_GIFT_MAX_QUEUE:]

    return {
        "status": "acknowledged",
        "giftId": normalized_gift_id,
        "remaining": len(next_queue)
    }


@app.get("/api/tea-gifts/receipts/{friend_id}")
def get_tea_gift_receipts(friend_id: str) -> Dict[str, Union[List[Dict[str, str]], int]]:
    normalized_friend_id = normalize_tea_gift_id(friend_id, "好友 ID")

    with tea_gift_lock:
        cleanup_expired_tea_gifts(datetime.now(timezone.utc))
        items = list(tea_gift_receipts.get(normalized_friend_id, []))

    return {
        "count": len(items),
        "items": items
    }


@app.delete("/api/tea-gifts/receipts/{friend_id}/{receipt_id}")
def acknowledge_tea_gift_receipt(friend_id: str, receipt_id: str) -> Dict[str, Union[str, int]]:
    normalized_friend_id = normalize_tea_gift_id(friend_id, "好友 ID")
    normalized_receipt_id = receipt_id.strip()

    if not normalized_receipt_id:
        raise HTTPException(status_code=400, detail="无效的奶茶回执。")

    with tea_gift_lock:
        cleanup_expired_tea_gifts(datetime.now(timezone.utc))
        queue = tea_gift_receipts.get(normalized_friend_id, [])
        next_queue = [receipt for receipt in queue if receipt["id"] != normalized_receipt_id]

        if next_queue:
            tea_gift_receipts[normalized_friend_id] = next_queue
        else:
            tea_gift_receipts.pop(normalized_friend_id, None)

    return {
        "status": "acknowledged",
        "receiptId": normalized_receipt_id,
        "remaining": len(next_queue)
    }


def libreoffice_convert(input_path: Path, output_format: str, output_path: Path) -> None:
    """Generic LibreOffice conversion."""
    import shutil
    import subprocess

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice 未安装，无法执行转换。")

    try:
        subprocess.run(
            [
                soffice,
                "--headless",
                "--norestore",
                "--convert-to", output_format,
                "--outdir", str(output_path.parent),
                str(input_path),
            ],
            capture_output=True,
            timeout=180,
            check=False,
        )
    except subprocess.TimeoutExpired:
        raise RuntimeError("转换超时，请尝试更小的文件。")

    # LibreOffice outputs with input stem + new extension
    ext = output_format.split(":")[0] if ":" in output_format else output_format
    lo_output = output_path.parent / f"{input_path.stem}.{ext}"
    if lo_output.exists() and lo_output.stat().st_size > 0:
        if lo_output != output_path:
            lo_output.rename(output_path)
    else:
        raise RuntimeError(f"LibreOffice 转换失败，未生成 .{ext} 文件。")


@app.post("/api/word-to-pdf")
async def word_to_pdf(file: UploadFile = File(...)) -> Response:
    fname = (file.filename or "").lower()
    if not (fname.endswith(".doc") or fname.endswith(".docx")):
        raise HTTPException(status_code=400, detail="请上传 Word 文件（.doc/.docx）。")

    with TemporaryDirectory(prefix="niuma-w2p-") as tmp:
        workspace = Path(tmp)
        input_path = workspace / (file.filename or "input.docx")
        output_path = workspace / f"{input_path.stem}.pdf"

        await save_upload(file, input_path)

        try:
            await run_in_threadpool(libreoffice_convert, input_path, "pdf", output_path)
        except Exception as exc:
            raise HTTPException(status_code=500, detail=f"转换失败：{exc}") from exc

        content = output_path.read_bytes()
        out_name = quote(f"{input_path.stem}.pdf")
        headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{out_name}"}
        return Response(content=content, media_type="application/pdf", headers=headers)


@app.post("/api/excel-to-csv")
async def excel_to_csv(
    file: UploadFile = File(...),
    encoding: str = Form("UTF-8"),
) -> Response:
    fname = (file.filename or "").lower()
    if not (fname.endswith(".xls") or fname.endswith(".xlsx")):
        raise HTTPException(status_code=400, detail="请上传 Excel 文件（.xls/.xlsx）。")

    with TemporaryDirectory(prefix="niuma-e2c-") as tmp:
        workspace = Path(tmp)
        input_path = workspace / (file.filename or "input.xlsx")
        output_path = workspace / f"{input_path.stem}.csv"

        await save_upload(file, input_path)

        try:
            await run_in_threadpool(
                libreoffice_convert, input_path, "csv:Text - txt - csv (StarCalc):44,34,76,1", output_path
            )
        except Exception as exc:
            raise HTTPException(status_code=500, detail=f"转换失败：{exc}") from exc

        # Re-encode if GBK requested
        if encoding.upper() == "GBK":
            text = output_path.read_text(encoding="utf-8", errors="replace")
            output_path.write_bytes(text.encode("gbk", errors="replace"))

        content = output_path.read_bytes()
        out_name = quote(f"{input_path.stem}.csv")
        headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{out_name}"}
        return Response(content=content, media_type="text/csv", headers=headers)


@app.post("/api/csv-to-excel")
async def csv_to_excel(file: UploadFile = File(...)) -> Response:
    fname = (file.filename or "").lower()
    if not fname.endswith(".csv"):
        raise HTTPException(status_code=400, detail="请上传 CSV 文件。")

    with TemporaryDirectory(prefix="niuma-c2e-") as tmp:
        workspace = Path(tmp)
        input_path = workspace / (file.filename or "input.csv")
        output_path = workspace / f"{input_path.stem}.xlsx"

        await save_upload(file, input_path)

        try:
            await run_in_threadpool(libreoffice_convert, input_path, "xlsx", output_path)
        except Exception as exc:
            raise HTTPException(status_code=500, detail=f"转换失败：{exc}") from exc

        content = output_path.read_bytes()
        xlsx_mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        out_name = quote(f"{input_path.stem}.xlsx")
        headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{out_name}"}
        return Response(content=content, media_type=xlsx_mime, headers=headers)


@app.post("/api/excel-merge")
async def excel_merge(files: list[UploadFile] = File(...)) -> Response:
    from openpyxl import Workbook, load_workbook

    valid_files = [upload for upload in files if (upload.filename or "").lower().endswith(".xlsx")]
    if len(valid_files) < 2:
        raise HTTPException(status_code=400, detail="请上传至少 2 个 .xlsx 文件。")

    with TemporaryDirectory(prefix="niuma-xmerge-") as tmp:
        workspace = Path(tmp)
        output_path = workspace / "merged.xlsx"
        output_book = Workbook()
        output_book.remove(output_book.active)
        used_titles: set[str] = set()

        for upload in valid_files:
            input_path = workspace / (upload.filename or f"input-{id(upload)}.xlsx")
            await save_upload(upload, input_path)
            source_book = load_workbook(input_path, data_only=False)

            for source_sheet in source_book.worksheets:
                title = safe_sheet_title(f"{input_path.stem}_{source_sheet.title}", used_titles)
                target_sheet = output_book.create_sheet(title=title)

                for row in source_sheet.iter_rows():
                    for cell in row:
                        target_sheet.cell(row=cell.row, column=cell.column, value=cell.value)

        output_book.save(output_path)
        content = output_path.read_bytes()
        headers = {"Content-Disposition": "attachment; filename*=UTF-8''merged.xlsx"}
        return Response(content=content, media_type=XLSX_MIME, headers=headers)


@app.post("/api/excel-split")
async def excel_split(file: UploadFile = File(...)) -> Response:
    from openpyxl import Workbook, load_workbook

    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="请上传 .xlsx 文件。")

    with TemporaryDirectory(prefix="niuma-xsplit-") as tmp:
        workspace = Path(tmp)
        input_path = workspace / (file.filename or "input.xlsx")
        zip_path = workspace / "split-sheets.zip"
        await save_upload(file, input_path)

        source_book = load_workbook(input_path, data_only=False)
        stem = Path(file.filename or "excel").stem

        with ZipFile(zip_path, "w", ZIP_DEFLATED) as archive:
            used_names: set[str] = set()
            for source_sheet in source_book.worksheets:
                target_book = Workbook()
                target_sheet = target_book.active
                target_sheet.title = safe_sheet_title(source_sheet.title, set())

                for row in source_sheet.iter_rows():
                    for cell in row:
                        target_sheet.cell(row=cell.row, column=cell.column, value=cell.value)

                filename = safe_sheet_title(f"{stem}_{source_sheet.title}", used_names) + ".xlsx"
                sheet_path = workspace / filename
                target_book.save(sheet_path)
                archive.write(sheet_path, filename)

        content = zip_path.read_bytes()
        out_name = quote(f"{stem}_sheets.zip")
        headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{out_name}"}
        return Response(content=content, media_type="application/zip", headers=headers)


def rows_from_json_data(data) -> list[list[object]]:
    if isinstance(data, list) and all(isinstance(item, dict) for item in data):
        headers = []
        for item in data:
            for key in item:
                if key not in headers:
                    headers.append(key)

        rows = [headers]
        for item in data:
            rows.append([
                json.dumps(item.get(key), ensure_ascii=False) if isinstance(item.get(key), (dict, list)) else item.get(key)
                for key in headers
            ])
        return rows

    if isinstance(data, dict):
        return [["key", "value"], *[
            [key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value]
            for key, value in data.items()
        ]]

    if isinstance(data, list):
        return [["value"], *[[json.dumps(item, ensure_ascii=False) if isinstance(item, (dict, list)) else item] for item in data]]

    return [["value"], [data]]


@app.post("/api/json-excel")
async def json_excel(file: UploadFile = File(...)) -> Response:
    from openpyxl import Workbook, load_workbook

    fname = (file.filename or "").lower()
    if not (fname.endswith(".json") or fname.endswith(".xlsx")):
        raise HTTPException(status_code=400, detail="请上传 JSON 或 XLSX 文件。")

    with TemporaryDirectory(prefix="niuma-json-excel-") as tmp:
        workspace = Path(tmp)
        input_path = workspace / (file.filename or "input")
        await save_upload(file, input_path)

        if fname.endswith(".json"):
            try:
                data = json.loads(input_path.read_text(encoding="utf-8"))
            except json.JSONDecodeError as exc:
                raise HTTPException(status_code=400, detail=f"JSON 格式无效：第 {exc.lineno} 行第 {exc.colno} 列。") from exc

            output_path = workspace / f"{input_path.stem}.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "JSON"

            for row in rows_from_json_data(data):
                sheet.append(row)

            book.save(output_path)
            content = output_path.read_bytes()
            out_name = quote(f"{input_path.stem}.xlsx")
            headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{out_name}"}
            return Response(content=content, media_type=XLSX_MIME, headers=headers)

        book = load_workbook(input_path, data_only=True)
        payload = {}
        for sheet in book.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                payload[sheet.title] = []
                continue

            headers = [str(value) if value is not None else f"column_{index + 1}" for index, value in enumerate(rows[0])]
            payload[sheet.title] = [
                {headers[index]: value for index, value in enumerate(row)}
                for row in rows[1:]
            ]

        output_path = workspace / f"{input_path.stem}.json"
        output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        content = output_path.read_bytes()
        out_name = quote(f"{input_path.stem}.json")
        headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{out_name}"}
        return Response(content=content, media_type="application/json", headers=headers)


@app.post("/api/pdf-merge")
async def pdf_merge(files: list[UploadFile] = File(...)) -> Response:
    if len(files) < 2:
        raise HTTPException(status_code=400, detail="请上传至少 2 个 PDF 文件。")

    with TemporaryDirectory(prefix="niuma-merge-") as tmp:
        workspace = Path(tmp)
        output_path = workspace / "merged.pdf"

        merged = fitz.open()
        try:
            for upload in files:
                if not (upload.filename or "").lower().endswith(".pdf"):
                    continue
                file_path = workspace / (upload.filename or f"file-{id(upload)}.pdf")
                await save_upload(upload, file_path)
                doc = fitz.open(str(file_path))
                merged.insert_pdf(doc)
                doc.close()

            if merged.page_count == 0:
                raise HTTPException(status_code=400, detail="没有有效的 PDF 页面可合并。")

            merged.save(str(output_path))
        finally:
            merged.close()

        content = output_path.read_bytes()
        headers = {"Content-Disposition": "attachment; filename*=UTF-8''merged.pdf"}
        return Response(content=content, media_type="application/pdf", headers=headers)


@app.post("/api/pdf-split")
async def pdf_split(
    file: UploadFile = File(...),
    start: int = Form(1),
    end: int = Form(-1),
) -> Response:
    if not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="请上传 PDF 文件。")

    with TemporaryDirectory(prefix="niuma-split-") as tmp:
        workspace = Path(tmp)
        input_path = workspace / "input.pdf"
        output_path = workspace / "split.pdf"

        await save_upload(file, input_path)

        pdf = fitz.open(str(input_path))
        try:
            total = pdf.page_count
            s = max(0, start - 1)
            e = total if end == -1 else min(end, total)

            if s >= e:
                raise HTTPException(status_code=400, detail="页码范围无效。")

            split_doc = fitz.open()
            split_doc.insert_pdf(pdf, from_page=s, to_page=e - 1)
            split_doc.save(str(output_path))
            split_doc.close()
        finally:
            pdf.close()

        content = output_path.read_bytes()
        stem = Path(file.filename or "split").stem
        out_name = quote(f"{stem}_p{start}-{end if end != -1 else 'end'}.pdf")
        headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{out_name}"}
        return Response(content=content, media_type="application/pdf", headers=headers)


@app.post("/api/pdf-compress")
async def pdf_compress(
    file: UploadFile = File(...),
    level: str = Form("标准压缩"),
) -> Response:
    if not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="请上传 PDF 文件。")

    with TemporaryDirectory(prefix="niuma-pdf-compress-") as tmp:
        workspace = Path(tmp)
        input_path = workspace / "input.pdf"
        output_path = workspace / "compressed.pdf"
        await save_upload(file, input_path)

        pdf = fitz.open(str(input_path))
        try:
            if pdf.needs_pass:
                raise HTTPException(status_code=400, detail="加密 PDF 请先使用 PDF 去密码后再压缩。")

            effort = 9 if level == "强力压缩" else 6
            pdf.save(
                str(output_path),
                garbage=4,
                clean=True,
                deflate=True,
                deflate_images=True,
                deflate_fonts=True,
                use_objstms=1,
                compression_effort=effort,
            )
        finally:
            pdf.close()

        content = output_path.read_bytes()
        stem = Path(file.filename or "compressed").stem
        out_name = quote(f"{stem}_compressed.pdf")
        headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{out_name}"}
        return Response(content=content, media_type="application/pdf", headers=headers)


@app.post("/api/pdf-organize")
async def pdf_organize(
    file: UploadFile = File(...),
    pages: str = Form("1-3"),
) -> Response:
    if not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="请上传 PDF 文件。")

    with TemporaryDirectory(prefix="niuma-pdf-organize-") as tmp:
        workspace = Path(tmp)
        input_path = workspace / "input.pdf"
        output_path = workspace / "organized.pdf"
        await save_upload(file, input_path)

        pdf = fitz.open(str(input_path))
        organized = fitz.open()
        try:
            if pdf.needs_pass:
                raise HTTPException(status_code=400, detail="加密 PDF 请先使用 PDF 去密码后再整理页面。")

            page_order = parse_pdf_page_order(pages, pdf.page_count)
            for page_index in page_order:
                organized.insert_pdf(pdf, from_page=page_index, to_page=page_index)

            organized.save(str(output_path), garbage=4, deflate=True)
        finally:
            organized.close()
            pdf.close()

        content = output_path.read_bytes()
        stem = Path(file.filename or "organized").stem
        out_name = quote(f"{stem}_organized.pdf")
        headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{out_name}"}
        return Response(content=content, media_type="application/pdf", headers=headers)


@app.post("/api/pdf-to-image")
async def pdf_to_image(
    file: UploadFile = File(...),
    image_type: str = Form("png"),
    page_range: str = Form("all"),
) -> Response:
    if not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="请上传 PDF 文件。")

    with TemporaryDirectory(prefix="niuma-p2i-") as tmp:
        workspace = Path(tmp)
        input_path = workspace / "input.pdf"
        await save_upload(file, input_path)

        pdf = fitz.open(str(input_path))
        try:
            total = pdf.page_count
            if page_range == "first":
                pages = [0]
            else:
                pages = list(range(total))

            image_paths = []
            for i in pages:
                page = pdf.load_page(i)
                pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0), alpha=False)
                ext = "png" if image_type.lower() == "png" else "jpg"
                img_path = workspace / f"page-{i + 1}.{ext}"
                if ext == "jpg":
                    pix.save(str(img_path), output="jpeg")
                else:
                    pix.save(str(img_path))
                image_paths.append(img_path)
        finally:
            pdf.close()

        if len(image_paths) == 1:
            content = image_paths[0].read_bytes()
            mime = "image/png" if image_type.lower() == "png" else "image/jpeg"
            ext = "png" if image_type.lower() == "png" else "jpg"
            stem = Path(file.filename or "page").stem
            out_name = quote(f"{stem}.{ext}")
            headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{out_name}"}
            return Response(content=content, media_type=mime, headers=headers)

        # Multiple pages -> zip
        zip_path = workspace / "pages.zip"
        with ZipFile(zip_path, "w", ZIP_DEFLATED) as zf:
            for img_path in image_paths:
                zf.write(img_path, img_path.name)

        content = zip_path.read_bytes()
        stem = Path(file.filename or "pages").stem
        out_name = quote(f"{stem}_images.zip")
        headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{out_name}"}
        return Response(content=content, media_type="application/zip", headers=headers)


@app.post("/api/pdf-unlock")
async def pdf_unlock(
    file: UploadFile = File(...),
    password: str = Form(""),
) -> Response:
    if not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="请上传 PDF 文件。")

    with TemporaryDirectory(prefix="niuma-unlock-") as tmp:
        workspace = Path(tmp)
        input_path = workspace / "input.pdf"
        output_path = workspace / "unlocked.pdf"
        await save_upload(file, input_path)

        pdf = fitz.open(str(input_path))
        try:
            if pdf.needs_pass:
                if not password:
                    raise HTTPException(status_code=400, detail="请输入 PDF 打开密码。")

                if not pdf.authenticate(password):
                    raise HTTPException(status_code=400, detail="PDF 密码不正确，请检查后重试。")

            pdf.save(
                str(output_path),
                encryption=fitz.PDF_ENCRYPT_NONE,
                garbage=4,
                deflate=True,
            )
        finally:
            pdf.close()

        content = output_path.read_bytes()
        stem = Path(file.filename or "unlocked").stem
        out_name = quote(f"{stem}_unlocked.pdf")
        headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{out_name}"}
        return Response(content=content, media_type="application/pdf", headers=headers)


@app.post("/api/image-ocr")
async def image_ocr(
    file: UploadFile = File(...),
    language: str = Form("chi_sim"),
) -> Response:
    import json
    import subprocess
    import shutil

    fname = (file.filename or "").lower()
    valid_exts = (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff")
    if not any(fname.endswith(ext) for ext in valid_exts):
        raise HTTPException(status_code=400, detail="请上传图片文件。")

    tesseract = shutil.which("tesseract")
    if not tesseract:
        raise HTTPException(status_code=500, detail="OCR 服务未安装（tesseract 未找到），请联系管理员。")

    with TemporaryDirectory(prefix="niuma-ocr-") as tmp:
        workspace = Path(tmp)
        input_path = workspace / (file.filename or "input.png")
        output_base = workspace / "result"

        await save_upload(file, input_path)

        # Map language option
        lang_map = {"中文优先": "chi_sim", "中英混排": "chi_sim+eng"}
        lang = lang_map.get(language, language)

        try:
            result = await run_in_threadpool(
                subprocess.run,
                [tesseract, str(input_path), str(output_base), "-l", lang, "--psm", "6"],
                capture_output=True,
                timeout=60,
            )
        except subprocess.TimeoutExpired:
            raise HTTPException(status_code=500, detail="OCR 识别超时，请尝试更小的图片。")

        output_file = workspace / "result.txt"
        if not output_file.exists():
            stderr = result.stderr.decode("utf-8", errors="replace") if result else ""
            raise HTTPException(status_code=500, detail=f"OCR 识别失败：{stderr[:200]}")

        text = output_file.read_text(encoding="utf-8").strip()

        return Response(
            content=json.dumps({"text": text}, ensure_ascii=False),
            media_type="application/json",
        )


@app.post("/api/pdf-watermark")
async def pdf_watermark(
    file: UploadFile = File(...),
    text: str = Form("仅供内部使用"),
    position: str = Form("居中斜排"),
) -> Response:
    if not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="请上传 PDF 文件。")

    with TemporaryDirectory(prefix="niuma-wm-") as tmp:
        workspace = Path(tmp)
        input_path = workspace / "input.pdf"
        output_path = workspace / "watermarked.pdf"
        await save_upload(file, input_path)

        pdf = fitz.open(str(input_path))
        try:
            for page in pdf:
                rect = page.rect
                fontsize = min(rect.width, rect.height) * 0.06

                if position == "页脚":
                    point = fitz.Point(rect.width / 2, rect.height - 30)
                    rotate = 0
                elif position == "右上角":
                    point = fitz.Point(rect.width - 20, 40)
                    rotate = 0
                else:  # 居中斜排
                    point = fitz.Point(rect.width / 2, rect.height / 2)
                    rotate = -45

                page.insert_text(
                    point,
                    text,
                    fontsize=fontsize,
                    fontname="china-s",
                    color=(0.6, 0.6, 0.6),
                    rotate=rotate,
                    overlay=True,
                )

            pdf.save(str(output_path))
        finally:
            pdf.close()

        content = output_path.read_bytes()
        stem = Path(file.filename or "watermarked").stem
        out_name = quote(f"{stem}_watermarked.pdf")
        headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{out_name}"}
        return Response(content=content, media_type="application/pdf", headers=headers)


@app.post("/api/pdf-to-word")
async def pdf_to_word(
    file: UploadFile = File(...),
    mode: str = Form(MODE_EDITABLE)
) -> Response:
    if mode not in SUPPORTED_MODES:
        raise HTTPException(status_code=400, detail="未知转换方案。")

    if not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="请上传 PDF 文件。")

    with TemporaryDirectory(prefix="niuma-pdf-") as tmp:
        workspace = Path(tmp)
        input_path = workspace / "input.pdf"
        output_name = safe_docx_name(file.filename)
        output_path = workspace / output_name

        await save_upload(file, input_path)

        try:
            await run_in_threadpool(convert_pdf, input_path, output_path, mode, workspace)
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=500, detail=f"转换失败：{exc}") from exc

        if not output_path.exists() or output_path.stat().st_size == 0:
            raise HTTPException(status_code=500, detail="转换失败：没有生成有效 Word 文件。")

        content = output_path.read_bytes()
        encoded_name = quote(output_name)
        headers = {
            "Content-Disposition": f"attachment; filename=\"converted.docx\"; filename*=UTF-8''{encoded_name}"
        }
        return Response(content=content, media_type=DOCX_MIME, headers=headers)
